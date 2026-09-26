from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


OUT = Path(__file__).with_name('evidence-template.xlsx')
HEADERS = [
    'CaseId', 'LogFileName', 'Environment', 'Server',
    'RunNumber', 'TargetDate', 'LogText',
]

book = Workbook()
sheet = book.active
sheet.title = '証跡'
sheet.sheet_view.showGridLines = False
sheet['A2'] = '架空ログの証跡（方式検証用）'
sheet['A2'].font = Font(name='Arial', size=15, bold=True, color='172B4D')

for col, title in enumerate(HEADERS, 1):
    cell = sheet.cell(row=4, column=col, value=title)
    cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='244C6B')
    cell.alignment = Alignment(vertical='center')

sheet['A5'] = '__TEMPLATE__'

for col, width in zip('ABCDEFG', [22, 24, 23, 21, 15, 18, 54]):
    sheet.column_dimensions[col].width = width
sheet.row_dimensions[4].height = 25

table = Table(displayName='Evidence', ref='A4:G5')
table._initialise_columns()
for column, title in zip(table.tableColumns, HEADERS):
    column.name = title
table.tableStyleInfo = TableStyleInfo(
    name='TableStyleMedium2', showFirstColumn=False,
    showLastColumn=False, showRowStripes=True, showColumnStripes=False,
)
sheet.add_table(table)
book.save(OUT)
print(OUT)
