"""Build the proposed one-case, multi-row evidence workbook.

The initial row is replaced with the first log chunk by the planned flow.
The remaining chunks are appended. Cloud compatibility is not established.
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).with_name("evidence-template.xlsx")
HEADERS = [
    "CaseId",
    "LogFileName",
    "Environment",
    "Server",
    "RunNumber",
    "TargetDate",
    "ChunkIndex",
    "ChunkCount",
    "LogTextPart",
]


def main() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "証跡"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    sheet["A2"] = "架空ログの証跡"
    sheet["A2"].font = Font(name="Arial", size=15, bold=True, color="172B4D")

    for column, heading in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=4, column=column, value=heading)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244C6B")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet["A5"] = "__TEMPLATE__"
    sheet["A5"].font = Font(name="Arial", size=10, color="6B7280")
    sheet.row_dimensions[4].height = 25

    for column, width in zip(
        "ABCDEFGHI", [38, 25, 22, 22, 14, 18, 15, 15, 70], strict=True
    ):
        sheet.column_dimensions[column].width = width
    for column in "GHI":
        sheet[f"{column}5"].alignment = Alignment(vertical="top", wrap_text=True)

    table = Table(displayName="Evidence", ref="A4:I5")
    table._initialise_columns()
    for table_column, heading in zip(table.tableColumns, HEADERS, strict=True):
        table_column.name = heading
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    book.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=False, data_only=False)
    check_sheet = check["証跡"]
    assert list(check_sheet.tables) == ["Evidence"]
    assert check_sheet.tables["Evidence"].ref == "A4:I5"
    assert [check_sheet.cell(4, n).value for n in range(1, 10)] == HEADERS
    assert check_sheet["A5"].value == "__TEMPLATE__"
    assert all(check_sheet.cell(5, n).value is None for n in range(2, 10))
    print(f"Created and checked: {OUTPUT}")


if __name__ == "__main__":
    main()
