"""Offline one-case transfer prototype; it never calls Power Platform."""

from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
import tempfile
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table


ROOT = Path(__file__).resolve().parents[2]
MAX_UTF16_UNITS = 30_000
HEADERS = [
    "CaseId",
    "LogFileName",
    "Environment",
    "Server",
    "RunNumber",
    "TargetDate",
    "ChunkIndex",
    "ChunkCount",
    "LogTextPart",
]


class StopCase(Exception):
    """A validation condition that should stop this case before file creation."""


def resolve_path(value: str | Path, base: Path = ROOT) -> Path:
    path = Path(value)
    return path if path.is_absolute() else base / path


def chunk_utf16(text: str, limit: int = MAX_UTF16_UNITS) -> list[str]:
    """Split by UTF-16 code units without cutting a Unicode surrogate pair."""
    if limit < 2:
        raise ValueError("limit must be at least 2 UTF-16 code units")
    chunks: list[str] = []
    current: list[str] = []
    units = 0
    for char in text:
        width = 2 if ord(char) > 0xFFFF else 1
        if units + width > limit and current:
            chunks.append("".join(current))
            current = []
            units = 0
        current.append(char)
        units += width
    if current or not chunks:
        chunks.append("".join(current))
    return chunks


def parse_log(text: str, case: dict[str, Any]) -> None:
    lines = text.splitlines()
    if not lines or not lines[0].startswith("処理日: "):
        raise StopCase("冒頭の処理日行がありません")

    date_text = lines[0][len("処理日: ") :]
    try:
        log_date = dt.date.fromisoformat(date_text)
    except ValueError as exc:
        raise StopCase("冒頭の処理日が有効なYYYY-MM-DDではありません") from exc
    if log_date.isoformat() != date_text:
        raise StopCase("冒頭の処理日がYYYY-MM-DD形式ではありません")
    if date_text != case["targetDate"]:
        raise StopCase("ログの処理日と対象処理日が一致しません")

    parsed: dict[str, str] = {}
    for line in lines[:5]:
        if ": " in line:
            key, value = line.split(": ", 1)
            parsed[key] = value
    for log_key, case_key, label in (
        ("環境", "environment", "環境"),
        ("サーバー", "server", "サーバー"),
        ("実行回", "runNumber", "実行回"),
    ):
        value = parsed.get(log_key)
        if not value:
            raise StopCase(f"ログの{label}がありません")
        if str(value) != str(case[case_key]):
            raise StopCase(f"ログの{label}が保存済みの値と一致しません")


def validate_case(case: dict[str, Any], log_text: str, log_path: Path) -> None:
    try:
        uuid.UUID(str(case["caseId"]))
    except (KeyError, ValueError, TypeError) as exc:
        raise StopCase("件IDがGUID形式ではありません") from exc

    if case.get("saved") is not True:
        raise StopCase("保存済みの件ではありません")
    if case.get("startStatus") != "開始受付済み":
        raise StopCase("処理開始が開始受付済みではありません")
    for key, label in (
        ("environment", "環境"),
        ("server", "サーバー"),
        ("targetDate", "対象処理日"),
        ("logFileName", "元ファイル名"),
    ):
        if not str(case.get(key, "")).strip():
            raise StopCase(f"{label}がありません")

    try:
        run_number = int(case["runNumber"])
    except (KeyError, TypeError, ValueError) as exc:
        raise StopCase("実行回が1以上の整数ではありません") from exc
    if run_number < 1 or str(run_number) != str(case["runNumber"]):
        raise StopCase("実行回が1以上の整数ではありません")
    if log_path.name != case["logFileName"]:
        raise StopCase("指定された元ファイル名と実ファイル名が一致しません")
    parse_log(log_text, case)


def validate_template(template: Path) -> None:
    if not template.is_file():
        raise StopCase("ひな形Excelが見つかりません")
    try:
        workbook = load_workbook(template, read_only=False, data_only=False)
    except Exception as exc:  # openpyxl raises several format-specific errors
        raise StopCase("ひな形Excelを開けません") from exc
    try:
        if "証跡" not in workbook.sheetnames:
            raise StopCase("ひな形に証跡シートがありません")
        sheet = workbook["証跡"]
        if "Evidence" not in sheet.tables:
            raise StopCase("ひな形にEvidence表がありません")
        table = sheet.tables["Evidence"]
        if table.ref != "A4:I5":
            raise StopCase("Evidence表の想定範囲がA4:I5ではありません")
        if [sheet.cell(4, col).value for col in range(1, 10)] != HEADERS:
            raise StopCase("Evidence表の列名が想定と一致しません")
        if sheet["A5"].value != "__TEMPLATE__" or any(
            sheet.cell(5, col).value is not None for col in range(2, 10)
        ):
            raise StopCase("Evidence表の仮行が想定と一致しません")
    finally:
        workbook.close()


def row_values(case: dict[str, Any], index: int, count: int, part: str) -> list[Any]:
    return [
        str(case["caseId"]).lower(),
        case["logFileName"],
        case["environment"],
        case["server"],
        int(case["runNumber"]),
        case["targetDate"],
        index,
        count,
        part,
    ]


def write_and_verify(
    case: dict[str, Any], log_text: str, template: Path, output_path: Path
) -> dict[str, Any]:
    parts = chunk_utf16(log_text)
    shutil.copy2(template, output_path)

    workbook = load_workbook(output_path, read_only=False, data_only=False)
    try:
        sheet = workbook["証跡"]
        table: Table = sheet.tables["Evidence"]
        for index, part in enumerate(parts, start=1):
            values = row_values(case, index, len(parts), part)
            row = 4 + index
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=value)
        table.ref = f"A4:I{4 + len(parts)}"
        workbook.save(output_path)
    finally:
        workbook.close()

    check = load_workbook(output_path, read_only=False, data_only=False)
    try:
        sheet = check["証跡"]
        table = sheet.tables["Evidence"]
        expected_ref = f"A4:I{4 + len(parts)}"
        if table.ref != expected_ref:
            raise RuntimeError("保存後のEvidence表の範囲が一致しません")
        rows = [
            [sheet.cell(row, col).value for col in range(1, 10)]
            for row in range(5, 5 + len(parts))
        ]
        expected = [
            row_values(case, index, len(parts), part)
            for index, part in enumerate(parts, start=1)
        ]
        if rows != expected:
            raise RuntimeError("保存後のExcel行が期待値と一致しません")
        reconstructed = "".join(str(row[8]) for row in rows)
        if reconstructed != log_text:
            raise RuntimeError("保存後の全文が元ログと一致しません")
    finally:
        check.close()
    return {"chunks": len(parts), "utf16_units": len(log_text.encode("utf-16-le")) // 2}


def run_case(case: dict[str, Any], template: Path, output_dir: Path) -> dict[str, Any]:
    case_id = str(case.get("caseId", ""))
    output_name = f"{case_id.lower()}.xlsx"
    output_path = output_dir / output_name
    copied = False
    try:
        log_path = resolve_path(case["logPath"])
        template_path = resolve_path(template)
        if not log_path.is_file():
            raise StopCase("元ログファイルが見つかりません")
        try:
            log_text = log_path.read_bytes().decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise StopCase("元ログをUTF-8として読めません") from exc
        validate_case(case, log_text, log_path)
        validate_template(template_path)
        output_dir.mkdir(parents=True, exist_ok=True)
        if output_path.exists():
            return {
                "caseId": case_id,
                "processingStatus": "結果不明",
                "excelCheckStatus": "未確認",
                "reason": "同じ件IDの出力Excelが既にあります。上書き・再転記しません。",
                "localWorkbook": output_name,
            }
        result = write_and_verify(case, log_text, template_path, output_path)
        copied = True
        return {
            "caseId": case_id,
            "processingStatus": "転記済み",
            "excelCheckStatus": "全文一致",
            "reason": None,
            "localWorkbook": output_name,
            **result,
        }
    except StopCase as exc:
        return {
            "caseId": case_id,
            "processingStatus": "停止",
            "excelCheckStatus": "未確認",
            "reason": str(exc),
            "localWorkbook": output_name if output_path.exists() else None,
        }
    except Exception as exc:
        # Preserve any output created before failure; a later run must inspect it.
        copied = copied or output_path.exists()
        return {
            "caseId": case_id,
            "processingStatus": "結果不明" if copied else "停止",
            "excelCheckStatus": "読出不能" if copied else "未確認",
            "reason": f"ローカル処理中に確認できないエラーが発生しました: {type(exc).__name__}",
            "localWorkbook": output_name if output_path.exists() else None,
        }


def case_for_sample(name: str, *, target_date: str | None = None) -> dict[str, Any]:
    path = ROOT / "samples" / "t006" / name
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    values: dict[str, str] = {}
    for line in lines[:5]:
        if ": " in line:
            key, value = line.split(": ", 1)
            values[key] = value
    return {
        "caseId": str(uuid.uuid4()),
        "saved": True,
        "startStatus": "開始受付済み",
        "environment": values.get("環境", "架空環境H"),
        "server": values.get("サーバー", "架空サーバーH"),
        "runNumber": values.get("実行回", "1"),
        "targetDate": target_date or values.get("処理日", "2026-09-25"),
        "logFileName": path.name,
        "logPath": str(path),
    }


def self_test() -> None:
    template = ROOT / "outputs" / "t006-20260925" / "evidence-template.xlsx"
    with tempfile.TemporaryDirectory(prefix="case-transfer-") as directory:
        output_dir = Path(directory) / "output"

        normal = case_for_sample("normal-a.txt")
        success = run_case(normal, template, output_dir)
        assert success["processingStatus"] == "転記済み", success
        assert success["excelCheckStatus"] == "全文一致", success
        assert success["chunks"] == 1, success

        duplicate = run_case(normal, template, output_dir)
        assert duplicate["processingStatus"] == "結果不明", duplicate
        assert "既にあります" in duplicate["reason"], duplicate

        long_case = case_for_sample("long-unicode.txt")
        long_result = run_case(long_case, template, output_dir)
        assert long_result["processingStatus"] == "転記済み", long_result
        assert long_result["chunks"] > 1, long_result

        stop_cases = [
            case_for_sample("date-mismatch.txt", target_date="2026-09-25"),
            case_for_sample("invalid-date.txt"),
            case_for_sample("missing-date.txt"),
            case_for_sample("missing-identifier.txt"),
        ]
        for case in stop_cases:
            result = run_case(case, template, output_dir)
            assert result["processingStatus"] == "停止", result
            assert result["localWorkbook"] is None, result

    print("Offline self-test passed: normal, duplicate guard, long Unicode, and four stop cases.")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--case-json", type=Path)
    parser.add_argument("--template", default="outputs/t006-20260925/evidence-template.xlsx")
    parser.add_argument("--output-dir", default="outputs/flow-implementation/local-output")
    args = parser.parse_args()

    if args.self_test:
        self_test()
        return
    if args.case_json is None:
        parser.error("use --self-test or provide --case-json")
    case_file = args.case_json if args.case_json.is_absolute() else ROOT / args.case_json
    case = json.loads(case_file.read_text(encoding="utf-8"))
    report = run_case(case, resolve_path(args.template), resolve_path(args.output_dir))
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
