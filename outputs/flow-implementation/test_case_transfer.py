import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from case_transfer import HEADERS, ROOT, case_for_sample, run_case


class CaseTransferTests(unittest.TestCase):
    def test_runless_and_legacy_logs_transfer_and_read_back_case_id_and_received_at(self):
        template = ROOT / "outputs" / "t006-20260925" / "evidence-template-received-at-jst.xlsx"
        for sample_name in ("normal-runless.txt", "normal-a.txt"):
            with self.subTest(sample=sample_name), tempfile.TemporaryDirectory() as directory:
                output_dir = Path(directory) / "output"
                case = case_for_sample(sample_name)
                case["createdOnUtc"] = "2026-09-27T03:04:05Z"

                result = run_case(case, template, output_dir)

                self.assertEqual(result["processingStatus"], "転記済み", result)
                self.assertEqual(result["excelCheckStatus"], "全文一致", result)
                workbook = load_workbook(output_dir / result["localWorkbook"], data_only=False)
                try:
                    sheet = workbook["証跡"]
                    self.assertEqual([sheet.cell(4, n).value for n in range(1, 10)], HEADERS)
                    row = [sheet.cell(5, n).value for n in range(1, 10)]
                    self.assertEqual(row[0], case["caseId"].lower())
                    self.assertEqual(row[4], "2026-09-27 12:04:05")
                    self.assertEqual(row[8], Path(case["logPath"]).read_text(encoding="utf-8-sig"))
                    if sample_name == "normal-a.txt":
                        self.assertIn("実行回: 1", row[8])
                finally:
                    workbook.close()

    def test_date_environment_and_server_mismatches_stop_before_workbook_creation(self):
        template = ROOT / "outputs" / "t006-20260925" / "evidence-template-received-at-jst.xlsx"
        mismatch_cases = []
        date_case = case_for_sample("normal-runless.txt", target_date="2026-09-26")
        mismatch_cases.append(date_case)
        environment_case = case_for_sample("normal-runless.txt")
        environment_case["environment"] = "架空環境の不一致"
        mismatch_cases.append(environment_case)
        server_case = case_for_sample("normal-runless.txt")
        server_case["server"] = "架空サーバーの不一致"
        mismatch_cases.append(server_case)

        with tempfile.TemporaryDirectory() as directory:
            output_dir = Path(directory) / "output"
            for case in mismatch_cases:
                with self.subTest(case=case):
                    result = run_case(case, template, output_dir)
                    self.assertEqual(result["processingStatus"], "停止", result)
                    self.assertIsNone(result["localWorkbook"], result)


if __name__ == "__main__":
    unittest.main()
